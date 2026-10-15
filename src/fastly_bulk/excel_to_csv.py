from __future__ import annotations

import csv
import logging
from collections import defaultdict
from pathlib import Path
from typing import Iterable, Iterator

import openpyxl

from .models import CSV_COLUMNS, ConvertResult, OriginRow, RejectedRow
from .validation import is_ip, is_valid_address, is_valid_hostname, parse_port, registrable_domain

log = logging.getLogger(__name__)

COL_FQDN = "fqdn_header_route"
COL_ORIGIN = "origin_server"
COL_PORT = "origin_port outbound f5"
COL_HOST_REWRITE = "host_rewrite"
COL_ENV = "ambiente F5"
COL_LB = "load_balancer f5"
COL_ROUTE = "route_num"
REQUIRED_COLUMNS = [COL_FQDN, COL_ORIGIN, COL_PORT, COL_HOST_REWRITE, COL_ENV, COL_LB, COL_ROUTE]

REJECTED_COLUMNS = [
    "excel_row", "reason", "fqdn_header_route", "origin_server",
    "origin_port outbound f5", "environment", "source_lb", "source_route",
]


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_excel(path: Path, sheet: str = "Todos") -> Iterator[tuple[int, dict[str, str]]]:
    """Yield (excel_row_number, {header: text}) for each data row of the sheet."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Sheet {sheet!r} not found. Available: {wb.sheetnames}")
        rows = wb[sheet].iter_rows(values_only=True)
        headers = [_text(h) for h in next(rows)]
        missing = [c for c in REQUIRED_COLUMNS if c not in headers]
        if missing:
            raise ValueError(f"Missing required columns in sheet {sheet!r}: {missing}")
        for idx, row in enumerate(rows, start=2):
            record = {h: _text(v) for h, v in zip(headers, row)}
            if any(record.values()):
                yield idx, record
    finally:
        wb.close()


def _split_domain(fqdn: str) -> str:
    host = fqdn.strip().lower()
    if ":" in host:
        head, _, tail = host.rpartition(":")
        if tail.isdigit():
            host = head
    return host


def convert_rows(raw_rows: Iterable[tuple[int, dict[str, str]]]) -> ConvertResult:
    result = ConvertResult()
    seen: dict[tuple[str, str, int], str] = {}
    rewrites_by_domain: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))

    for excel_row, rec in raw_rows:
        result.rows_read += 1
        fqdn_raw, origin_raw, port_raw = rec[COL_FQDN], rec[COL_ORIGIN], rec[COL_PORT]

        def reject(reason: str) -> None:
            result.rejected.append(RejectedRow(
                excel_row, reason, fqdn_raw, origin_raw, port_raw,
                rec[COL_ENV], rec[COL_LB], rec[COL_ROUTE],
            ))

        if not fqdn_raw:
            reject("empty fqdn_header_route")
            continue
        if not origin_raw:
            reject("empty origin_server")
            continue
        if not port_raw:
            reject("empty origin_port")
            continue

        domain = _split_domain(fqdn_raw)
        if not is_valid_hostname(domain):
            reject(f"invalid domain: {domain!r}")
            continue
        address = origin_raw if is_ip(origin_raw) else origin_raw.lower()
        if not is_valid_address(address):
            reject(f"invalid origin address: {origin_raw!r}")
            continue
        port = parse_port(port_raw)
        if port is None:
            reject(f"invalid origin port: {port_raw!r}")
            continue

        host_rewrite = rec[COL_HOST_REWRITE]
        key = (domain, address, port)
        if key in seen:
            if seen[key] == host_rewrite:
                result.duplicates_dropped += 1
                continue
            result.warnings.append(
                f"row {excel_row}: same (domain, origin, port) {key} already present with a "
                f"different host_rewrite ({seen[key]!r} vs {host_rewrite!r}); keeping both"
            )
        else:
            seen[key] = host_rewrite

        rewrites_by_domain[domain][host_rewrite].append(f"row {excel_row}")
        result.rows.append(OriginRow(
            service_name=domain, domain=domain, origin_address=address, origin_port=port,
            host_rewrite=host_rewrite, environment=rec[COL_ENV],
            source_lb=rec[COL_LB], source_route=rec[COL_ROUTE],
        ))

        if host_rewrite:
            hr = host_rewrite.lower()
            if not is_valid_hostname(hr):
                result.warnings.append(
                    f"row {excel_row}: host_rewrite {host_rewrite!r} for {domain} is not a plain hostname"
                )
            elif registrable_domain(hr) != registrable_domain(domain):
                result.warnings.append(
                    f"row {excel_row}: host_rewrite {host_rewrite!r} has a different registrable "
                    f"domain than {domain!r} (possible typo); NOT corrected"
                )

    for domain, variants in rewrites_by_domain.items():
        if len(variants) > 1:
            detail = "; ".join(f"{hr or '<empty>'!r} ({', '.join(rows)})" for hr, rows in variants.items())
            result.warnings.append(f"domain {domain} has different host_rewrite values: {detail}")

    return result


def write_csv(result: ConvertResult, csv_path: Path, rejected_path: Path) -> None:
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    rejected_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=CSV_COLUMNS)
        writer.writeheader()
        for row in result.rows:
            writer.writerow(row.as_csv_dict())
    with rejected_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(REJECTED_COLUMNS)
        for r in result.rejected:
            writer.writerow([
                r.excel_row, r.reason, r.raw_fqdn, r.raw_origin_server, r.raw_origin_port,
                r.environment, r.source_lb, r.source_route,
            ])


def summary(result: ConvertResult) -> str:
    lines = [
        "=== convert summary ===",
        f"rows read:          {result.rows_read}",
        f"rows written:       {len(result.rows)}",
        f"duplicates dropped: {result.duplicates_dropped}",
        f"unique domains:     {result.unique_domains}",
        f"unique origins:     {result.unique_origins}",
        f"rejected rows:      {len(result.rejected)}",
        f"warnings:           {len(result.warnings)}",
    ]
    if result.rejected:
        lines.append("")
        lines.append("rejected:")
        lines += [f"  row {r.excel_row}: {r.reason}" for r in result.rejected]
    if result.warnings:
        lines.append("")
        lines.append("warnings:")
        lines += [f"  - {w}" for w in result.warnings]
    return "\n".join(lines)


def run_convert(input_path: Path, sheet: str, csv_path: Path, rejected_path: Path) -> ConvertResult:
    result = convert_rows(read_excel(input_path, sheet))
    write_csv(result, csv_path, rejected_path)
    return result
